import argparse
import collections
import html
import math
import os
import random
import re
from dataclasses import dataclass
from typing import Dict, Iterable, List, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import PatternFill
from scipy import sparse
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import linear_kernel
from sklearn.decomposition import TruncatedSVD
from sklearn.preprocessing import normalize


EXCEL_MAX_ROWS = 1_000_000
RANDOM_SEED = 42
STOPWORDS = {
    "the", "and", "for", "with", "from", "that", "this", "have", "has", "are", "you", "your", "will", "our",
    "job", "role", "work", "working", "experience", "candidate", "position", "company", "looking", "required",
    "skills", "good", "full", "time", "part", "offer", "office", "use", "using", "ability", "level", "new",
    "team", "within", "knowledge", "years", "year", "english", "italian", "available", "activities", "support",
}
TITLE_FILL_COLORS = [
    "FFF2CC",
    "D9EAD3",
    "D0E0E3",
    "F4CCCC",
    "D9D2E9",
    "FCE5CD",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a candidate-job matching ranking for validation data."
    )
    parser.add_argument(
        "--train",
        required=True,
        help="Path to Matching_training.xlsx",
    )
    parser.add_argument(
        "--validation",
        required=True,
        help="Path to Matching_validation.xlsx",
    )
    parser.add_argument(
        "--output",
        default="matching_validation_scored.xlsx",
        help="Output workbook path.",
    )
    parser.add_argument(
        "--top-n",
        type=int,
        default=20,
        help="How many top jobs per candidate to include in the summary sheet.",
    )
    parser.add_argument(
        "--candidate-batch-size",
        type=int,
        default=64,
        help="Candidates processed per similarity batch.",
    )
    return parser.parse_args()


def clean_text(value: object) -> str:
    text = "" if pd.isna(value) else str(value)
    text = html.unescape(text)
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("\n", " ").replace("\r", " ").replace("\t", " ")
    text = re.sub(r"[^a-zA-Z0-9%€$+#./ -]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip().lower()


def normalize_title(value: object) -> str:
    return clean_text(value)


def build_job_document(title: object, body: object) -> str:
    title_text = normalize_title(title)
    body_text = clean_text(body)
    repeated_title = " ".join([title_text] * 5) if title_text else ""
    return " ".join(part for part in [repeated_title, body_text] if part).strip()


def tokenize(text: str) -> List[str]:
    tokens = re.findall(r"[a-zA-Z][a-zA-Z0-9+#./-]{1,}", text.lower())
    return [token for token in tokens if token not in STOPWORDS and len(token) > 2]


def summarize_overlap(candidate_text: str, job_title: str, job_text: str, top_k: int = 5) -> str:
    candidate_tokens = tokenize(candidate_text)
    title_tokens = tokenize(job_title)
    job_tokens = tokenize(job_text)

    candidate_counter = collections.Counter(candidate_tokens)
    title_counter = collections.Counter(title_tokens)
    job_counter = collections.Counter(job_tokens)
    overlap_scores: Dict[str, int] = {}

    for token in set(candidate_counter) & set(job_counter):
        overlap_scores[token] = candidate_counter[token] + job_counter[token] + 2 * title_counter[token]

    if not overlap_scores:
        return "The candidate profile shows a general semantic alignment with the responsibilities and skills described in this role."

    ranked_terms = sorted(overlap_scores.items(), key=lambda item: (-item[1], item[0]))[:top_k]
    important_terms = [term for term, _ in ranked_terms]

    title_hits = [term for term in important_terms if term in title_counter]
    if title_hits:
        supporting_terms = [term for term in important_terms if term not in title_hits][: max(0, top_k - len(title_hits))]
        return (
            f"The candidate appears well aligned with the role '{job_title.strip()}' because the CV highlights experience related to "
            + ", ".join(title_hits + supporting_terms)
            + "."
        )

    return (
        "The match is mainly supported by shared experience and skill signals in the CV and job description, especially around "
        + ", ".join(important_terms)
        + "."
    )


@dataclass
class DatasetBundle:
    candidates: pd.DataFrame
    jobs: pd.DataFrame
    matches: pd.DataFrame | None = None


def load_training_dataset(path: str) -> DatasetBundle:
    candidates = pd.read_excel(path, sheet_name="Candidates")
    jobs = pd.read_excel(path, sheet_name="Job_posts")
    matches = pd.read_excel(path, sheet_name="Matches")
    return DatasetBundle(candidates=candidates, jobs=jobs, matches=matches)


def load_validation_dataset(path: str) -> DatasetBundle:
    candidates = pd.read_excel(path, sheet_name="Candidates")
    jobs = pd.read_excel(path, sheet_name="Job_posts")
    return DatasetBundle(candidates=candidates, jobs=jobs)


def prepare_candidates(df: pd.DataFrame) -> pd.DataFrame:
    prepared = df.copy()
    prepared["cv_clean"] = prepared["CV_text"].map(clean_text)
    return prepared


def prepare_jobs(df: pd.DataFrame) -> pd.DataFrame:
    prepared = df.copy()
    prepared["job_title_clean"] = prepared["Job_title_eng"].map(normalize_title)
    prepared["job_post_clean"] = prepared["Job_post_eng"].map(clean_text)
    prepared["job_document"] = [
        build_job_document(title, body)
        for title, body in zip(prepared["Job_title_eng"], prepared["Job_post_eng"])
    ]
    return prepared


def fit_vectorizers(all_candidate_docs: Sequence[str], all_job_docs: Sequence[str], all_titles: Sequence[str]):
    corpus_main = list(all_candidate_docs) + list(all_job_docs)
    corpus_titles = list(all_candidate_docs) + list(all_titles)

    word_vectorizer = TfidfVectorizer(
        ngram_range=(1, 2),
        min_df=2,
        max_features=60000,
        strip_accents="unicode",
        sublinear_tf=True,
    )
    char_vectorizer = TfidfVectorizer(
        analyzer="char_wb",
        ngram_range=(3, 5),
        min_df=2,
        max_features=40000,
        sublinear_tf=True,
    )
    title_vectorizer = TfidfVectorizer(
        ngram_range=(1, 2),
        min_df=2,
        max_features=30000,
        strip_accents="unicode",
        sublinear_tf=True,
    )

    word_vectorizer.fit(corpus_main)
    char_vectorizer.fit(corpus_main)
    title_vectorizer.fit(corpus_titles)
    return word_vectorizer, char_vectorizer, title_vectorizer


def fit_svd_embeddings(
    train_candidate_word: sparse.csr_matrix,
    train_job_word: sparse.csr_matrix,
    train_candidate_title: sparse.csr_matrix,
    train_job_title: sparse.csr_matrix,
    n_components_main: int = 256,
    n_components_title: int = 128,
):
    svd_main = TruncatedSVD(n_components=n_components_main, random_state=RANDOM_SEED)
    svd_title = TruncatedSVD(n_components=n_components_title, random_state=RANDOM_SEED)

    svd_main.fit(sparse.vstack([train_candidate_word, train_job_word]))
    svd_title.fit(sparse.vstack([train_candidate_title, train_job_title]))
    return svd_main, svd_title


def transform_embeddings(svd: TruncatedSVD, matrix: sparse.csr_matrix) -> np.ndarray:
    return normalize(svd.transform(matrix))


def transform_texts(vectorizer: TfidfVectorizer, texts: Sequence[str]) -> sparse.csr_matrix:
    return vectorizer.transform(list(texts)).tocsr()


def sparse_row_dot(matrix_a: sparse.csr_matrix, idx_a: int, matrix_b: sparse.csr_matrix, idx_b: int) -> float:
    return float(matrix_a[idx_a].multiply(matrix_b[idx_b]).sum())


def tune_rank_weights(
    matches: pd.DataFrame,
    candidate_ids: Sequence[int],
    job_ids: Sequence[str],
    feature_matrices: Sequence[np.ndarray],
) -> np.ndarray:
    candidate_index = {int(candidate_id): idx for idx, candidate_id in enumerate(candidate_ids)}
    job_ids_arr = np.array(list(map(str, job_ids)), dtype=object)

    ranked_matches = matches.sort_values(["CANDIDATE_ID", "match_score"], ascending=[True, False])
    holdout = ranked_matches.groupby("CANDIDATE_ID", group_keys=False).head(1)
    ground_truth = {
        candidate_index[int(row.CANDIDATE_ID)]: str(row.Offer_ID)
        for row in holdout.itertuples(index=False)
        if int(row.CANDIDATE_ID) in candidate_index
    }

    trial_weights: List[Tuple[float, ...]] = [
        (2.03, 2.68, 0.26, 1.27, 0.09),
        (0.66, 2.99, 1.53, 0.27, 0.14),
        (1.19, 2.74, 1.38, 0.79, 0.74),
        (0.33, 1.88, 2.38, 1.27, 0.19),
        (1.00, 1.00, 1.50, 1.00, 0.50),
        (1.50, 2.50, 1.00, 1.00, 0.25),
    ]
    rng = random.Random(RANDOM_SEED)
    for _ in range(32):
        trial_weights.append(tuple(round(rng.uniform(0.0, 3.0), 2) for _ in range(5)))

    best_metric = (-1.0, -1.0, -1.0, -1.0)
    best_weights = np.array(trial_weights[0], dtype=np.float32)
    for weights in trial_weights:
        score_matrix = sum(weight * feature for weight, feature in zip(weights, feature_matrices))
        recall_1 = 0
        recall_5 = 0
        recall_10 = 0
        reciprocal_rank = 0.0

        for candidate_idx, true_offer_id in ground_truth.items():
            ranking = np.argsort(-score_matrix[candidate_idx])
            ranked_offer_ids = job_ids_arr[ranking]
            true_position = int(np.where(ranked_offer_ids == true_offer_id)[0][0]) + 1
            recall_1 += int(true_position <= 1)
            recall_5 += int(true_position <= 5)
            recall_10 += int(true_position <= 10)
            reciprocal_rank += 1.0 / true_position

        total = max(1, len(ground_truth))
        metric = (
            recall_5 / total,
            recall_10 / total,
            reciprocal_rank / total,
            recall_1 / total,
        )
        if metric > best_metric:
            best_metric = metric
            best_weights = np.array(weights, dtype=np.float32)

    return best_weights


def score_batch(
    candidate_word_batch: sparse.csr_matrix,
    candidate_char_batch: sparse.csr_matrix,
    candidate_title_batch: sparse.csr_matrix,
    job_word: sparse.csr_matrix,
    job_char: sparse.csr_matrix,
    job_title: sparse.csr_matrix,
    candidate_main_emb_batch: np.ndarray,
    candidate_title_emb_batch: np.ndarray,
    job_main_emb: np.ndarray,
    job_title_emb: np.ndarray,
    weights: np.ndarray,
) -> np.ndarray:
    word_scores = linear_kernel(candidate_word_batch, job_word)
    char_scores = linear_kernel(candidate_char_batch, job_char)
    title_scores = linear_kernel(candidate_title_batch, job_title)
    main_emb_scores = np.clip(candidate_main_emb_batch @ job_main_emb.T, 0.0, 1.0)
    title_emb_scores = np.clip(candidate_title_emb_batch @ job_title_emb.T, 0.0, 1.0)

    stacked = np.stack([word_scores, char_scores, title_scores, main_emb_scores, title_emb_scores], axis=-1).astype(
        np.float32
    )
    predictions = np.tensordot(stacked, weights, axes=([-1], [0]))
    max_weight_sum = float(max(1.0, np.sum(weights)))
    predictions = np.clip(predictions / max_weight_sum, 0.0, 1.0)
    return predictions.astype(np.float32)


def write_dataframe_sheet(workbook: Workbook, name: str, df: pd.DataFrame) -> None:
    ws = workbook.create_sheet(title=name)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))


def write_colored_dataframe_sheet(workbook: Workbook, name: str, df: pd.DataFrame, group_column: str) -> None:
    ws = workbook.create_sheet(title=name)
    header = list(df.columns)
    ws.append(header)

    group_values = list(dict.fromkeys(df[group_column].astype(str).tolist()))
    fill_map = {
        group: PatternFill(fill_type="solid", fgColor=TITLE_FILL_COLORS[idx % len(TITLE_FILL_COLORS)])
        for idx, group in enumerate(group_values)
    }

    for row in df.itertuples(index=False, name=None):
        group_value = str(row[header.index(group_column)])
        fill = fill_map[group_value]
        styled_row = []
        for value in row:
            cell = WriteOnlyCell(ws, value=value)
            cell.fill = fill
            styled_row.append(cell)
        ws.append(styled_row)


def export_results(
    validation_candidates: pd.DataFrame,
    validation_jobs: pd.DataFrame,
    score_batches: Iterable[Tuple[pd.DataFrame, np.ndarray]],
    output_path: str,
    top_n: int,
) -> None:
    workbook = Workbook(write_only=True)
    top_5_rows: List[Dict[str, object]] = []

    write_dataframe_sheet(workbook, "Candidates", validation_candidates[["CANDIDATE_ID", "CV_text"]])
    write_dataframe_sheet(workbook, "Job_posts", validation_jobs[["Offer_ID", "Job_post_eng", "Job_title_eng"]])

    summary_sheet = workbook.create_sheet(title=f"Top_{top_n}_per_candidate")
    summary_sheet.append(
        [
            "CANDIDATE_ID",
            "rank_for_candidate",
            "Offer_ID",
            "affinity_score",
            "Job_title_eng",
            "match_reason",
        ]
    )

    sheet_idx = 1
    current_rows = 0
    result_sheet = workbook.create_sheet(title=f"Predictions_{sheet_idx}")
    header = ["CANDIDATE_ID", "Offer_ID", "rank_for_candidate", "affinity_score", "Job_title_eng"]
    result_sheet.append(header)
    current_rows += 1

    job_ids = validation_jobs["Offer_ID"].astype(str).to_numpy()
    job_titles = validation_jobs["Job_title_eng"].astype(str).to_numpy()
    job_posts = validation_jobs["Job_post_eng"].astype(str).to_numpy()
    candidate_text_map = dict(
        zip(validation_candidates["CANDIDATE_ID"].astype(int), validation_candidates["CV_text"].astype(str))
    )

    for candidate_batch_df, predictions in score_batches:
        candidate_ids = candidate_batch_df["CANDIDATE_ID"].astype(int).to_numpy()

        for local_idx, candidate_id in enumerate(candidate_ids):
            candidate_scores = predictions[local_idx]
            ranking = np.argsort(-candidate_scores)

            top_limit = min(top_n, len(ranking))
            for rank in range(top_limit):
                job_idx = int(ranking[rank])
                candidate_id_int = int(candidate_id)
                match_reason = summarize_overlap(
                    candidate_text=candidate_text_map[candidate_id_int],
                    job_title=str(job_titles[job_idx]),
                    job_text=str(job_posts[job_idx]),
                )
                summary_sheet.append(
                    [
                        candidate_id_int,
                        rank + 1,
                        str(job_ids[job_idx]),
                        float(round(candidate_scores[job_idx] * 100.0, 6)),
                        str(job_titles[job_idx]),
                        match_reason,
                    ]
                )
                if rank < 5:
                    top_5_rows.append(
                        {
                            "CANDIDATE_ID": candidate_id_int,
                            "rank_for_candidate": rank + 1,
                            "Offer_ID": str(job_ids[job_idx]),
                            "affinity_score": float(round(candidate_scores[job_idx] * 100.0, 6)),
                            "Job_title_eng": str(job_titles[job_idx]),
                            "match_reason": match_reason,
                        }
                    )

            for rank, job_idx in enumerate(ranking, start=1):
                if current_rows >= EXCEL_MAX_ROWS:
                    sheet_idx += 1
                    result_sheet = workbook.create_sheet(title=f"Predictions_{sheet_idx}")
                    result_sheet.append(header)
                    current_rows = 1

                result_sheet.append(
                    [
                        int(candidate_id),
                        str(job_ids[job_idx]),
                        rank,
                        float(round(candidate_scores[job_idx] * 100.0, 6)),
                        str(job_titles[job_idx]),
                    ]
                )
                current_rows += 1

    top_5_df = pd.DataFrame(top_5_rows)
    if not top_5_df.empty:
        write_dataframe_sheet(workbook, "Top_5_per_candidate", top_5_df)
        top_5_sorted_df = top_5_df.sort_values(
            by=["Job_title_eng", "affinity_score", "CANDIDATE_ID"],
            ascending=[True, False, True],
        ).reset_index(drop=True)
        write_colored_dataframe_sheet(
            workbook,
            "Top_5_sorted_by_job_title",
            top_5_sorted_df,
            group_column="Job_title_eng",
        )

    workbook.save(output_path)


def main() -> int:
    args = parse_args()
    random.seed(RANDOM_SEED)
    np.random.seed(RANDOM_SEED)

    train_path = os.path.abspath(args.train)
    validation_path = os.path.abspath(args.validation)
    output_path = os.path.abspath(args.output)

    training = load_training_dataset(train_path)
    validation = load_validation_dataset(validation_path)

    train_candidates = prepare_candidates(training.candidates)
    train_jobs = prepare_jobs(training.jobs)
    val_candidates = prepare_candidates(validation.candidates)
    val_jobs = prepare_jobs(validation.jobs)

    word_vectorizer, char_vectorizer, title_vectorizer = fit_vectorizers(
        all_candidate_docs=list(train_candidates["cv_clean"]) + list(val_candidates["cv_clean"]),
        all_job_docs=list(train_jobs["job_document"]) + list(val_jobs["job_document"]),
        all_titles=list(train_jobs["job_title_clean"]) + list(val_jobs["job_title_clean"]),
    )

    train_candidate_word = transform_texts(word_vectorizer, train_candidates["cv_clean"])
    train_candidate_char = transform_texts(char_vectorizer, train_candidates["cv_clean"])
    train_candidate_title = transform_texts(title_vectorizer, train_candidates["cv_clean"])
    train_job_word = transform_texts(word_vectorizer, train_jobs["job_document"])
    train_job_char = transform_texts(char_vectorizer, train_jobs["job_document"])
    train_job_title = transform_texts(title_vectorizer, train_jobs["job_title_clean"])
    svd_main, svd_title = fit_svd_embeddings(
        train_candidate_word=train_candidate_word,
        train_job_word=train_job_word,
        train_candidate_title=train_candidate_title,
        train_job_title=train_job_title,
    )
    train_candidate_main_emb = transform_embeddings(svd_main, train_candidate_word)
    train_job_main_emb = transform_embeddings(svd_main, train_job_word)
    train_candidate_title_emb = transform_embeddings(svd_title, train_candidate_title)
    train_job_title_emb = transform_embeddings(svd_title, train_job_title)

    val_candidate_word = transform_texts(word_vectorizer, val_candidates["cv_clean"])
    val_candidate_char = transform_texts(char_vectorizer, val_candidates["cv_clean"])
    val_candidate_title = transform_texts(title_vectorizer, val_candidates["cv_clean"])
    val_job_word = transform_texts(word_vectorizer, val_jobs["job_document"])
    val_job_char = transform_texts(char_vectorizer, val_jobs["job_document"])
    val_job_title = transform_texts(title_vectorizer, val_jobs["job_title_clean"])
    val_candidate_main_emb = transform_embeddings(svd_main, val_candidate_word)
    val_job_main_emb = transform_embeddings(svd_main, val_job_word)
    val_candidate_title_emb = transform_embeddings(svd_title, val_candidate_title)
    val_job_title_emb = transform_embeddings(svd_title, val_job_title)

    candidate_index = {
        int(candidate_id): idx for idx, candidate_id in enumerate(train_candidates["CANDIDATE_ID"].astype(int))
    }
    job_index = {str(offer_id): idx for idx, offer_id in enumerate(train_jobs["Offer_ID"].astype(str))}

    train_feature_matrices = [
        linear_kernel(train_candidate_word, train_job_word).astype(np.float32),
        linear_kernel(train_candidate_char, train_job_char).astype(np.float32),
        linear_kernel(train_candidate_title, train_job_title).astype(np.float32),
        np.clip(train_candidate_main_emb @ train_job_main_emb.T, 0.0, 1.0).astype(np.float32),
        np.clip(train_candidate_title_emb @ train_job_title_emb.T, 0.0, 1.0).astype(np.float32),
    ]
    weights = tune_rank_weights(
        matches=training.matches,
        candidate_ids=list(train_candidates["CANDIDATE_ID"].astype(int)),
        job_ids=list(train_jobs["Offer_ID"].astype(str)),
        feature_matrices=train_feature_matrices,
    )

    def score_batches():
        batch_size = max(1, int(args.candidate_batch_size))
        for start in range(0, len(val_candidates), batch_size):
            end = min(start + batch_size, len(val_candidates))
            candidate_batch_df = val_candidates.iloc[start:end][["CANDIDATE_ID"]]
            predictions = score_batch(
                val_candidate_word[start:end],
                val_candidate_char[start:end],
                val_candidate_title[start:end],
                val_job_word,
                val_job_char,
                val_job_title,
                val_candidate_main_emb[start:end],
                val_candidate_title_emb[start:end],
                val_job_main_emb,
                val_job_title_emb,
                weights,
            )
            yield candidate_batch_df, predictions

    export_results(
        validation_candidates=validation.candidates,
        validation_jobs=validation.jobs,
        score_batches=score_batches(),
        output_path=output_path,
        top_n=args.top_n,
    )

    total_pairs = len(val_candidates) * len(val_jobs)
    required_prediction_sheets = math.ceil(total_pairs / (EXCEL_MAX_ROWS - 1))
    print(f"Training candidates: {len(train_candidates)}")
    print(f"Training jobs: {len(train_jobs)}")
    print(f"Training matches: {len(training.matches)}")
    print(f"Validation candidates: {len(val_candidates)}")
    print(f"Validation jobs: {len(val_jobs)}")
    print(f"Validation pairs scored: {total_pairs}")
    print(f"Prediction sheets written: {required_prediction_sheets}")
    print(f"Output saved to: {output_path}")
    print(
        "Tuned weights: "
        f"word={weights[0]:.4f}, char={weights[1]:.4f}, title={weights[2]:.4f}, "
        f"main_emb={weights[3]:.4f}, title_emb={weights[4]:.4f}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
